#!/usr/bin/env python3
"""Erzeugt einen eigenständigen englischen Testtext für den Korrektur-UX Test.

Nur für Testdaten. Erzeugt eine DOCX-Datei mit einem englischen Essay
über Climate Change (7. Klasse = Unterstufe).

Mehrfach ausführbar (überschreibt bestehende Datei).
"""

from pathlib import Path

try:
    from docx import Document
except ImportError:
    raise SystemExit("python-docx muss installiert sein: pip install python-docx")


KLASSE = "TEST-7a"
DATEINAME = "benchmark_en_a2_essay.docx"


def main() -> int:
    doc = Document()
    doc.add_paragraph("Synthetic English Benchmark Case")
    doc.add_paragraph("")
    doc.add_paragraph("Climate Change and Its Effects on Our Daily Lives")
    doc.add_paragraph("")
    doc.add_paragraph(
        "Climate change is one of the biggest problems we face today. "
        "The Earth is getting warmer because of greenhouse gases like carbon dioxide. "
        "This leads to more extreme weather, such as floods and droughts."
    )
    doc.add_paragraph("")
    doc.add_paragraph(
        "In my opinion, everyone should try to help. "
        "We can use less plastic, ride our bikes more often, and save energy at home. "
        "My family has started recycling and we try to use public transport instead of the car."
    )
    doc.add_paragraph("")
    doc.add_paragraph(
        "However, it is also important that governments make new laws. "
        "For example, they could tax companies that produce too much CO2. "
        "I think this would help a lot to protect our planet for the future."
    )
    doc.add_paragraph("")
    doc.add_paragraph(
        "In conclusion, climate change is a serious issue, but if everyone does their part, "
        "we can make a difference. The future of our planet depends on what we do today."
    )

    output_dir = Path(__file__).resolve().parent / "input" / KLASSE
    output_dir.mkdir(parents=True, exist_ok=True)
    docx_path = output_dir / DATEINAME
    doc.save(str(docx_path))
    print(f"Englischer Testtext erstellt: {docx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
